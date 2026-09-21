"""Erzeugt den SAP-Massenupload aus abgeschlossenen Rückblicken."""

from __future__ import annotations

import hashlib
import json
import re
import sqlite3
import uuid
from copy import copy
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from .packages import bounded_review_period


HEADERS = [
    "PersNr",
    "Beurteilungsart",
    "Beginndatum IT9075",
    "Endedatum IT9075",
    "Ans.",
    "Datum MAB",
    "Beurteilungszeitraum von",
    "Beurteilungszeitraum bis",
    "Gesamtbeurteilung",
    "Zielerreichung",
    "Fachliche Kompetenz",
    "Sozialkompetenz (Verhalten)",
    "Führungskompetenz",
    "Nächster Termin",
]


@dataclass(frozen=True)
class SapUploadRow:
    personal_number: int
    assessment_type: int
    it9075_start: date
    it9075_end: date
    employment_assignment: int
    dialog_date: date
    assessment_period_start: date
    assessment_period_end: date
    overall_rating: str
    dialog_event_id: int = 0
    case_id: str = ""
    source_document_id: int | None = None
    source_version: str = ""

    def as_excel_row(self) -> list[Any]:
        return [
            self.personal_number,
            self.assessment_type,
            self.it9075_start,
            self.it9075_end,
            self.employment_assignment,
            self.dialog_date,
            self.assessment_period_start,
            self.assessment_period_end,
            self.overall_rating,
            None,
            None,
            None,
            None,
            None,
        ]


def _required_integer(value: Any, label: str, pn: str) -> int:
    text = str(value or "").strip()
    if not text or not text.isdigit():
        raise ValueError(f"PN {pn}: {label} fehlt oder ist nicht numerisch.")
    return int(text)


def _required_iso_date(value: Any, label: str, pn: str) -> date:
    text = str(value or "").strip()
    try:
        return date.fromisoformat(text)
    except ValueError as exc:
        raise ValueError(f"PN {pn}: {label} fehlt oder ist ungültig.") from exc


def _rating_code(value: Any, pn: str) -> str:
    match = re.match(r"^\s*([A-E])\b", str(value or ""), flags=re.IGNORECASE)
    if not match:
        raise ValueError(f"PN {pn}: Gesamtbeurteilung A bis E fehlt.")
    return match.group(1).upper()


def collect_sap_upload_rows(
    connection: sqlite3.Connection, cycle_id: int
) -> list[SapUploadRow]:
    cycle = connection.execute(
        "SELECT id, review_year FROM cycles WHERE id = ?", (cycle_id,)
    ).fetchone()
    if not cycle:
        raise LookupError("Jahresprozess nicht gefunden.")

    cases = connection.execute(
        """
        SELECT dc.case_id, dc.employee_pn, dc.data_json, dc.official_scope,
               dc.dialog_date, dc.overall_rating_code,
               e.entry_date, e.exit_date,
               COALESCE(NULLIF(dc.employment_assignment, ''), NULLIF(e.employment_assignment, ''), '1') AS employment_assignment,
               de.id AS dialog_event_id, COALESCE(de.sap_leading, 0) AS sap_leading,
               (
                   SELECT od.id FROM official_documents od
                   WHERE od.case_id = dc.case_id AND od.document_kind = 'review'
                     AND od.is_current = 1
                   ORDER BY CASE od.variant WHEN 'digital' THEN 0 ELSE 1 END, od.id DESC
                   LIMIT 1
               ) AS source_document_id,
               (
                   SELECT od.sha256 FROM official_documents od
                   WHERE od.case_id = dc.case_id AND od.document_kind = 'review'
                     AND od.is_current = 1
                   ORDER BY CASE od.variant WHEN 'digital' THEN 0 ELSE 1 END, od.id DESC
                   LIMIT 1
               ) AS source_document_sha
        FROM dialog_cases dc
        JOIN employees e ON e.pn = dc.employee_pn
        LEFT JOIN dialog_events de ON de.legacy_case_id = dc.case_id
        WHERE dc.cycle_id = ?
          AND dc.status = 'vollstaendig'
          AND NOT EXISTS (
              SELECT 1 FROM sap_export_rows ser WHERE ser.dialog_event_id = de.id
          )
        ORDER BY CAST(dc.employee_pn AS INTEGER), dc.employee_pn, dc.manager_pn
        """,
        (cycle_id,),
    ).fetchall()

    rows: list[SapUploadRow] = []
    review_candidates: list[tuple[sqlite3.Row, dict[str, Any], str]] = []
    for case in cases:
        pn = str(case["employee_pn"])
        payload: dict[str, Any] = {}
        if case["data_json"]:
            try:
                payload = json.loads(case["data_json"])
            except json.JSONDecodeError as exc:
                raise ValueError(f"PN {pn}: gespeicherter MD-Datensatz ist ungültig.") from exc
        scope = case["official_scope"] or payload.get("scope", "")
        if scope not in {"full", "review_only"}:
            continue
        review_candidates.append((case, payload, scope))

    leading_keys = {
        (str(case["employee_pn"]), str(case["employment_assignment"] or "").strip())
        for case, _payload, _scope in review_candidates if case["sap_leading"]
    }
    unresolved_keys = {
        (str(case["employee_pn"]), str(case["employment_assignment"] or "").strip())
        for case, _payload, _scope in review_candidates
        if not case["sap_leading"]
    } - leading_keys
    if unresolved_keys:
        pn, assignment = sorted(unresolved_keys)[0]
        raise ValueError(
            f"PN {pn}, Ans. {assignment or 'leer'}: Das führende SAP-Ereignis ist nicht festgelegt."
        )

    for case, payload, _scope in review_candidates:
        if not case["sap_leading"]:
            continue
        pn = str(case["employee_pn"])

        assignment = str(case["employment_assignment"] or "").strip()

        if not case["entry_date"]:
            raise ValueError(
                f"PN {pn}: Eintrittsdatum fehlt; der SAP-Zeitraum kann nicht sicher geprüft werden."
            )
        period_start_text, period_end_text = bounded_review_period(
            int(cycle["review_year"]), case["entry_date"], case["exit_date"]
        )
        period_start = date.fromisoformat(period_start_text)
        period_end = date.fromisoformat(period_end_text)
        review = payload.get("review") or {}
        dialog_date = case["dialog_date"] or payload.get("dialog_date")
        overall_rating = case["overall_rating_code"] or review.get("overall_rating")

        source_version = case["source_document_sha"] or hashlib.sha256(
            json.dumps(
                {
                    "case_id": case["case_id"],
                    "source_document_sha": case["source_document_sha"] or "",
                    "data_json": case["data_json"],
                    "official_scope": case["official_scope"],
                    "employment_assignment": assignment,
                    "dialog_date": dialog_date,
                    "overall_rating": overall_rating,
                    "period_start": period_start.isoformat(),
                    "period_end": period_end.isoformat(),
                },
                ensure_ascii=False,
                sort_keys=True,
            ).encode("utf-8")
        ).hexdigest()
        rows.append(
            SapUploadRow(
                personal_number=_required_integer(pn, "Personalnummer", pn),
                assessment_type=1,
                it9075_start=period_start,
                it9075_end=period_end,
                employment_assignment=_required_integer(assignment, "Ans.", pn),
                dialog_date=_required_iso_date(dialog_date, "Datum MAB", pn),
                assessment_period_start=period_start,
                assessment_period_end=period_end,
                overall_rating=_rating_code(overall_rating, pn),
                dialog_event_id=int(case["dialog_event_id"]),
                case_id=case["case_id"],
                source_document_id=case["source_document_id"],
                source_version=source_version,
            )
        )

    if not rows:
        raise ValueError("Es gibt keine neuen, noch nicht exportierten SAP-relevanten Rückblicke.")
    return rows


def write_sap_upload(
    rows: list[SapUploadRow], template_path: Path, output_path: Path
) -> Path:
    workbook = load_workbook(template_path)
    worksheet = workbook["Massenupload"] if "Massenupload" in workbook.sheetnames else workbook.active
    actual_headers = [worksheet.cell(1, column).value for column in range(1, len(HEADERS) + 1)]
    if actual_headers != HEADERS:
        raise ValueError("Die SAP-Massenupload-Vorlage hat nicht die erwartete Spaltenstruktur.")

    template_styles = []
    for column in range(1, len(HEADERS) + 1):
        cell = worksheet.cell(2, column)
        template_styles.append(
            {
                "font": copy(cell.font),
                "fill": copy(cell.fill),
                "border": copy(cell.border),
                "alignment": copy(cell.alignment),
                "protection": copy(cell.protection),
                "number_format": cell.number_format,
            }
        )
    if worksheet.max_row > 1:
        worksheet.delete_rows(2, worksheet.max_row - 1)

    for row_number, row in enumerate(rows, start=2):
        for column, value in enumerate(row.as_excel_row(), start=1):
            cell = worksheet.cell(row_number, column, value)
            style = template_styles[column - 1]
            cell.font = copy(style["font"])
            cell.fill = copy(style["fill"])
            cell.border = copy(style["border"])
            cell.alignment = copy(style["alignment"])
            cell.protection = copy(style["protection"])
            cell.number_format = "DD.MM.YYYY" if column in {3, 4, 6, 7, 8} else style["number_format"]

    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)
    return output_path


def create_sap_upload_file(
    connection: sqlite3.Connection,
    *,
    cycle_id: int,
    template_path: Path,
    output_dir: Path,
    created_at: datetime | None = None,
) -> Path:
    return create_sap_upload_batch(
        connection,
        cycle_id=cycle_id,
        template_path=template_path,
        output_dir=output_dir,
        created_at=created_at,
    )["path"]


def create_sap_upload_batch(
    connection: sqlite3.Connection,
    *,
    cycle_id: int,
    template_path: Path,
    output_dir: Path,
    user_id: int | None = None,
    created_at: datetime | None = None,
) -> dict[str, Any]:
    cycle = connection.execute(
        "SELECT review_year FROM cycles WHERE id = ?", (cycle_id,)
    ).fetchone()
    if not cycle:
        raise LookupError("Jahresprozess nicht gefunden.")
    created = created_at or datetime.now().astimezone()
    timestamp = created.strftime("%Y%m%d_%H%M%S")
    token = uuid.uuid4().hex[:8]
    output_path = output_dir / (
        f"SAP_Massenupload_MD_{cycle['review_year']}_{timestamp}_{token}.xlsx"
    )
    rows = collect_sap_upload_rows(connection, cycle_id)
    write_sap_upload(rows, template_path, output_path)
    file_hash = hashlib.sha256(output_path.read_bytes()).hexdigest()
    try:
        cursor = connection.execute(
            """
            INSERT INTO sap_export_batches (
                cycle_id, filename, stored_path, sha256, row_count, user_id, created_at
            ) VALUES (?, ?, ?, ?, ?, ?, ?)
            """,
            (
                cycle_id, output_path.name, str(output_path), file_hash, len(rows), user_id,
                created.isoformat(timespec="seconds"),
            ),
        )
        batch_id = cursor.lastrowid
        for row in rows:
            connection.execute(
                """
                INSERT INTO sap_export_rows (
                    batch_id, dialog_event_id, case_id, source_document_id,
                    source_version, employee_pn, employment_assignment,
                    assessment_type, it9075_start, it9075_end, dialog_date,
                    assessment_period_start, assessment_period_end, overall_rating
                ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                """,
                (
                    batch_id, row.dialog_event_id, row.case_id, row.source_document_id,
                    row.source_version, str(row.personal_number),
                    str(row.employment_assignment), row.assessment_type,
                    row.it9075_start.isoformat(), row.it9075_end.isoformat(),
                    row.dialog_date.isoformat(), row.assessment_period_start.isoformat(),
                    row.assessment_period_end.isoformat(), row.overall_rating,
                ),
            )
        connection.commit()
    except Exception:
        connection.rollback()
        output_path.unlink(missing_ok=True)
        raise
    return {
        "id": batch_id,
        "path": output_path,
        "filename": output_path.name,
        "sha256": file_hash,
        "row_count": len(rows),
    }


def sap_export_batches(
    connection: sqlite3.Connection, *, cycle_id: int | None = None
) -> list[sqlite3.Row]:
    where = "WHERE seb.cycle_id = ?" if cycle_id is not None else ""
    params = (cycle_id,) if cycle_id is not None else ()
    return connection.execute(
        f"""
        SELECT seb.*, c.review_year, c.outlook_year,
               COALESCE(u.email, 'System/Testbetrieb') AS user_email
        FROM sap_export_batches seb
        JOIN cycles c ON c.id = seb.cycle_id
        LEFT JOIN app_users u ON u.id = seb.user_id
        {where}
        ORDER BY seb.created_at DESC, seb.id DESC
        """,
        params,
    ).fetchall()


def get_sap_export_batch(connection: sqlite3.Connection, batch_id: int) -> sqlite3.Row:
    row = connection.execute(
        "SELECT * FROM sap_export_batches WHERE id = ?", (batch_id,)
    ).fetchone()
    if not row:
        raise LookupError("SAP-Exportbatch nicht gefunden.")
    if not Path(row["stored_path"]).is_file():
        raise ValueError("Die Datei dieses Exportbatches wurde lokal nicht gefunden.")
    return row
