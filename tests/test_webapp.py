from __future__ import annotations

import io
import json
import tempfile
import unittest
import zipfile
from datetime import date, datetime, timezone
from pathlib import Path
from xml.sax.saxutils import escape

from openpyxl import load_workbook
from pypdf import PdfReader
from reportlab.pdfgen import canvas

from prototype.html_dialog.generate_package import render_html
from prototype.html_dialog.validate_package import extract_payload
from webapp import create_app
from webapp.db import close_db, connect_database, get_db
from webapp.adapters.outlook import (
    DraftResult,
    InboundAttachment,
    InboundMessage,
    _resolve_inbox,
)
from webapp.services.cycles import change_cycle_status, create_cycle, cycle_overview
from webapp.services.cockpit import cockpit_overview
from webapp.services.dialog_events import set_leading_sap_event
from webapp.services.documents import (
    confirm_digital_signature,
    import_handwritten_scan,
    import_official_pdf,
)
from webapp.services.mail_dispatch import create_outlook_drafts, dispatch_candidates
from webapp.services.mail_intake import (
    mail_inbox_overview,
    process_inbound_messages,
    resolve_inbound_message,
)
from webapp.services.legacy_forms import import_legacy_forms, scan_legacy_forms
from webapp.services.case_review import case_review_detail, correct_case_data
from webapp.services.deadlines import extend_deadlines
from webapp.services.reminders import (
    confirm_reminder_sent,
    create_reminder_drafts,
    reminder_candidates,
)
from webapp.services.analytics import analytics_data
from webapp.services.operations import (
    create_system_backup,
    restore_system_backup,
    validate_backup,
)
from webapp.services.packages import (
    build_manager_payload,
    create_package_file,
    create_update_file,
    import_returned_package,
    manager_package_state,
)
from webapp.services.sap_export import create_sap_upload_file
from webapp.services.sap_import import import_sap_workbook

from tests.sample_data import create_sap_sample


PROJECT_ROOT = Path(__file__).resolve().parents[1]


class FakeOutlookDraftAdapter:
    def __init__(self, *, encryption_flag_verified: bool = True) -> None:
        self.encryption_flag_verified = encryption_flag_verified
        self.calls: list[dict] = []

    def create_encrypted_draft(self, **message) -> DraftResult:
        self.calls.append(message)
        return DraftResult(
            entry_id=f"draft-{len(self.calls)}",
            encryption_flag_verified=self.encryption_flag_verified,
        )

    def send_encrypted(self, **message) -> DraftResult:
        self.calls.append(message)
        return DraftResult(
            entry_id=f"sent-{len(self.calls)}",
            encryption_flag_verified=self.encryption_flag_verified,
        )


def write_test_pdf(path: Path, data_block: str) -> None:
    document = canvas.Canvas(str(path))
    document.drawString(40, 800, "Mitarbeitenden-Dialog Test")
    document.drawString(40, 770, data_block)
    document.save()


def write_legacy_docx(path: Path, controls: list[tuple[str, str]]) -> None:
    content = "".join(
        "<w:sdt><w:sdtPr><w:tag w:val=\"{}\"/></w:sdtPr>"
        "<w:sdtContent><w:p><w:r><w:t>{}</w:t></w:r></w:p></w:sdtContent></w:sdt>".format(
            escape(tag, {'"': '&quot;'}), escape(value)
        )
        for tag, value in controls
    )
    xml = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
        f"<w:body>{content}</w:body></w:document>"
    )
    with zipfile.ZipFile(path, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("word/document.xml", xml)


class WebAppIntegrationTest(unittest.TestCase):
    def setUp(self) -> None:
        self.temp_dir = tempfile.TemporaryDirectory()
        root = Path(self.temp_dir.name)
        self.sap_sample = create_sap_sample(root / "EXPORT_TEST.xlsx")
        self.app = create_app(
            {
                "TESTING": True,
                "SECRET_KEY": "test-secret",
                "DATABASE": str(root / "test.sqlite3"),
                "STORAGE_ROOT": str(root / "data"),
                "BACKUP_ROOT": str(root / "backups"),
                "AUTH_DISABLED": True,
            }
        )
        self.client = self.app.test_client()

    def tearDown(self) -> None:
        self.temp_dir.cleanup()

    def _prepare_cycle(self) -> int:
        with self.app.app_context():
            connection = get_db()
            imported = import_sap_workbook(
                connection,
                self.sap_sample,
                original_filename="EXPORT.XLSX",
                stored_filename="test_EXPORT.XLSX",
                imported_at=datetime(2025, 9, 17, 12, 0, tzinfo=timezone.utc),
            )
            return create_cycle(
                connection,
                review_year=2025,
                sap_import_id=imported.import_id,
                created_at=datetime(2025, 9, 17, 12, 1, tzinfo=timezone.utc),
            )

    def test_dashboard_starts_with_empty_database(self) -> None:
        response = self.client.get("/")
        self.assertEqual(response.status_code, 200)
        self.assertIn("Mitarbeitenden-Dialoge verwalten", response.get_data(as_text=True))
        self.assertIn("SAP-Stammdaten importieren", response.get_data(as_text=True))
        self.assertEqual(response.headers["Cache-Control"], "no-store")
        self.assertEqual(response.headers["X-Frame-Options"], "DENY")

    def test_administration_creates_complete_valid_backup(self) -> None:
        marker = Path(self.app.config["STORAGE_ROOT"]) / "marker.txt"
        marker.write_text("vorher", encoding="utf-8")
        response = self.client.post("/administration/backups")
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.mimetype, "application/zip")
        response_data = response.get_data()
        response.close()
        with zipfile.ZipFile(io.BytesIO(response_data)) as archive:
            self.assertIn("manifest.json", archive.namelist())
            self.assertIn("database.sqlite3", archive.namelist())
            self.assertIn("storage/marker.txt", archive.namelist())
        page = self.client.get("/administration")
        self.assertEqual(page.status_code, 200)
        self.assertIn("Betrieb und Audit", page.get_data(as_text=True))

    def test_backup_restore_roundtrip_and_tamper_detection(self) -> None:
        database_path = Path(self.app.config["DATABASE"])
        storage_root = Path(self.app.config["STORAGE_ROOT"])
        dossier_root = Path(self.app.config["DOSSIER_HANDOFF_ROOT"])
        backup_root = Path(self.app.config["BACKUP_ROOT"])
        marker = storage_root / "marker.txt"
        marker.write_text("gesichert", encoding="utf-8")
        with self.app.app_context():
            connection = get_db()
            connection.execute(
                "INSERT INTO audit_log (action, details_json, created_at) VALUES ('vorher', '{}', '2026-09-21T10:00:00+00:00')"
            )
            connection.commit()
            backup = create_system_backup(
                connection,
                database_path=database_path,
                storage_root=storage_root,
                dossier_root=dossier_root,
                backup_root=backup_root,
            )
            connection.execute(
                "INSERT INTO audit_log (action, details_json, created_at) VALUES ('nachher', '{}', '2026-09-21T11:00:00+00:00')"
            )
            connection.commit()
            marker.write_text("verändert", encoding="utf-8")
            close_db()
            result = restore_system_backup(
                backup["path"],
                database_path=database_path,
                storage_root=storage_root,
                dossier_root=dossier_root,
            )
            self.assertGreaterEqual(result["file_count"], 2)

        restored = connect_database(database_path)
        try:
            actions = [row[0] for row in restored.execute("SELECT action FROM audit_log ORDER BY id")]
        finally:
            restored.close()
        self.assertIn("vorher", actions)
        self.assertNotIn("nachher", actions)
        self.assertEqual(marker.read_text(encoding="utf-8"), "gesichert")

        tampered = backup_root / "tampered.zip"
        with zipfile.ZipFile(backup["path"]) as source, zipfile.ZipFile(tampered, "w") as target:
            for info in source.infolist():
                content = source.read(info.filename)
                target.writestr(
                    info,
                    b"manipuliert" if info.filename == "database.sqlite3" else content,
                )
        with self.assertRaisesRegex(ValueError, "Prüfsumme"):
            validate_backup(tampered)

    def test_legacy_word_goals_are_previewed_imported_and_added_to_package(self) -> None:
        cycle_id = self._prepare_cycle()
        forms_root = Path(self.temp_dir.name) / "legacy_forms"
        forms_root.mkdir()
        with self.app.app_context():
            connection = get_db()
            case = connection.execute(
                """
                SELECT employee_pn, employment_assignment, manager_pn
                FROM dialog_cases WHERE cycle_id = ? ORDER BY id LIMIT 1
                """,
                (cycle_id,),
            ).fetchone()
            write_legacy_docx(
                forms_root / "Ausblick_Test.docx",
                [
                    ("ab_pn", case["employee_pn"]),
                    ("ab_ziel", "Serviceprozess vereinfachen"),
                    ("ab_ziel_kriterien", "Durchlaufzeit messbar reduzieren"),
                    ("ab_ziel_schritte", "Prozess aufnehmen und anpassen"),
                    ("ab_ziel_termin", "2025-12-31"),
                    ("ab_entwicklungsfaehigkeit_ziel", "Feedback systematisch nutzen"),
                    ("ab_entwicklungsfaehigkeit_kriterien", "Quartalsweise Rückmeldung einholen"),
                    ("ab_entwicklungsfaehigkeit_schritte", "Feedbacktermine vereinbaren"),
                    ("ab_entwicklungsfaehigkeit_termin", "2025-11-30"),
                ],
            )
            preview = scan_legacy_forms(
                connection, forms_root=forms_root, goal_year=2025
            )
            self.assertEqual(preview["counts"], {"matched": 1})
            self.assertEqual(preview["performance_goals"], 1)
            self.assertEqual(preview["development_goals"], 1)
            self.assertEqual(
                connection.execute("SELECT COUNT(*) FROM historical_goals").fetchone()[0], 0
            )

            imported = import_legacy_forms(
                connection, forms_root=forms_root, goal_year=2025,
                imported_at=datetime(2026, 9, 21, 12, 0, tzinfo=timezone.utc),
            )
            self.assertEqual(imported["imported"], 1)
            self.assertEqual(
                connection.execute("SELECT COUNT(*) FROM historical_goals").fetchone()[0], 2
            )
            payload = build_manager_payload(
                connection, cycle_id=cycle_id, manager_pn=case["manager_pn"]
            )
            employee = next(
                item for item in payload["employees"]
                if item["employee"]["pn"] == case["employee_pn"]
            )
            self.assertEqual(
                employee["previous_goals"][0]["title"], "Serviceprozess vereinfachen"
            )
            self.assertEqual(
                employee["previous_development_goals"][0]["competency"],
                "Entwicklungsfähigkeit",
            )
            repeated = scan_legacy_forms(
                connection, forms_root=forms_root, goal_year=2025
            )
            self.assertEqual(repeated["counts"], {"already_imported": 1})

    def test_legacy_import_skips_identical_files_within_same_batch(self) -> None:
        cycle_id = self._prepare_cycle()
        forms_root = Path(self.temp_dir.name) / "legacy_duplicates"
        forms_root.mkdir()
        duplicate_root = forms_root / "test"
        duplicate_root.mkdir()
        with self.app.app_context():
            connection = get_db()
            case = connection.execute(
                """
                SELECT employee_pn FROM dialog_cases
                WHERE cycle_id = ? ORDER BY id LIMIT 1
                """,
                (cycle_id,),
            ).fetchone()
            original = forms_root / "Ausblick_Original.docx"
            write_legacy_docx(
                original,
                [("ab_pn", case["employee_pn"]), ("ab_ziel", "Einmal importieren")],
            )
            duplicate = duplicate_root / "Ausblick_Kopie.docx"
            duplicate.write_bytes(original.read_bytes())

            preview = scan_legacy_forms(
                connection, forms_root=forms_root, goal_year=2025
            )
            self.assertEqual(
                preview["counts"], {"matched": 1, "duplicate_in_batch": 1}
            )
            duplicate_item = next(
                item for item in preview["items"]
                if item["status"] == "duplicate_in_batch"
            )
            self.assertEqual(duplicate_item["duplicate_of"], "Ausblick_Original.docx")

            imported = import_legacy_forms(
                connection, forms_root=forms_root, goal_year=2025
            )
            self.assertEqual(imported["imported"], 1)
            self.assertEqual(
                connection.execute("SELECT COUNT(*) FROM legacy_form_imports").fetchone()[0],
                1,
            )

    def test_main_navigation_and_module_pages(self) -> None:
        self._prepare_cycle()
        pages = {
            "/": "Übersicht",
            "/stammdaten": "SAP-Datenstände",
            "/dialoge": "Durchläufe und Dialogpflichten",
            "/versand": "Arbeitsmappen bereitstellen",
            "/ruecklaeufe": "PDFs prüfen und verarbeiten",
            "/sap-export": "Massenuploads erstellen",
            "/auswertungen": "Prozess- und Organisationsauswertungen",
        }
        for path, heading in pages.items():
            with self.subTest(path=path):
                response = self.client.get(path)
                self.assertEqual(response.status_code, 200)
                page = response.get_data(as_text=True)
                self.assertIn(heading, page)
                self.assertIn("Stammdaten", page)
                self.assertIn("Rückläufe", page)

    def test_operational_overview_counts_missing_workbooks(self) -> None:
        cycle_id = self._prepare_cycle()
        with self.app.app_context():
            data = cockpit_overview(get_db(), selected_cycle_id=cycle_id)
            self.assertEqual(data["metrics"]["workbooks_missing"], 3)
            self.assertEqual(data["metrics"]["cases_open"], 9)
            self.assertEqual(
                len([task for task in data["tasks"] if task["kind"] == "Arbeitsmappe"]),
                3,
            )

        package = self.client.post(f"/cycles/{cycle_id}/managers/111116/package")
        package.close()
        with self.app.app_context():
            data = cockpit_overview(get_db(), selected_cycle_id=cycle_id)
            self.assertEqual(data["metrics"]["workbooks_missing"], 2)

    def test_process_guide_and_controlled_cycle_phase_changes(self) -> None:
        cycle_id = self._prepare_cycle()
        dashboard = self.client.get("/").get_data(as_text=True)
        self.assertIn("Wo befindet sich dieser Durchlauf?", dashboard)
        self.assertIn("Phase 1 von 4", dashboard)
        self.assertIn("Regulären und unterjährigen Ablauf anzeigen", dashboard)
        self.assertIn("Was bewirkt die App ausserhalb des Cockpits?", dashboard)

        page = self.client.get(f"/cycles/{cycle_id}").get_data(as_text=True)
        self.assertIn("Versandphase starten", page)
        response = self.client.post(
            f"/cycles/{cycle_id}/status",
            data={"status": "versand"},
            follow_redirects=True,
        )
        self.assertIn("Phase «Versand»", response.get_data(as_text=True))
        self.assertIn("Rücklaufphase starten", response.get_data(as_text=True))
        with self.app.app_context():
            connection = get_db()
            self.assertEqual(
                connection.execute(
                    "SELECT status FROM cycles WHERE id = ?", (cycle_id,)
                ).fetchone()["status"],
                "versand",
            )
            audit_row = connection.execute(
                "SELECT action, details_json FROM audit_log ORDER BY id DESC LIMIT 1"
            ).fetchone()
            self.assertEqual(audit_row["action"], "cycle_status_changed")
            self.assertIn('"new_status": "versand"', audit_row["details_json"])

        response = self.client.post(
            f"/cycles/{cycle_id}/status",
            data={"status": "abgeschlossen"},
            follow_redirects=True,
        )
        self.assertIn(
            "nur um einen Schritt vor- oder zurückgesetzt", response.get_data(as_text=True)
        )

    def test_cycle_status_service_can_reopen_completed_cycle(self) -> None:
        cycle_id = self._prepare_cycle()
        with self.app.app_context():
            connection = get_db()
            for status in ("versand", "ruecklauf", "abgeschlossen", "ruecklauf"):
                change_cycle_status(
                    connection, cycle_id=cycle_id, new_status=status
                )
            current = connection.execute(
                "SELECT status FROM cycles WHERE id = ?", (cycle_id,)
            ).fetchone()["status"]
            self.assertEqual(current, "ruecklauf")

    def test_start_and_update_files_are_detected_and_generated(self) -> None:
        cycle_id = self._prepare_cycle()
        manager_page = self.client.get(
            f"/cycles/{cycle_id}/managers/111116"
        ).get_data(as_text=True)
        self.assertIn("data-refresh-after-download", manager_page)
        self.assertIn("dispatch.js", manager_page)
        with self.app.app_context(), tempfile.TemporaryDirectory() as temp:
            connection = get_db()
            create_package_file(
                connection, cycle_id=cycle_id, manager_pn="111116",
                output_dir=Path(temp) / "packages",
            )
            self.assertEqual(
                manager_package_state(
                    connection, cycle_id=cycle_id, manager_pn="111116"
                )["state"],
                "current",
            )
            connection.execute(
                "UPDATE employees SET org_unit = 'abt-neu' WHERE pn = '111111'"
            )
            connection.commit()
            self.assertEqual(
                manager_package_state(
                    connection, cycle_id=cycle_id, manager_pn="111116"
                )["state"],
                "update_required",
            )
            path, payload = create_update_file(
                connection, cycle_id=cycle_id, manager_pn="111116",
                output_dir=Path(temp) / "packages",
            )
            self.assertEqual(path.suffix, ".json")
            self.assertEqual(payload["schema_version"], "1.0-update")
            self.assertEqual(payload["update"]["manager_pn"], "111116")
            self.assertEqual(
                manager_package_state(
                    connection, cycle_id=cycle_id, manager_pn="111116"
                )["state"],
                "current",
            )

    def test_replacement_start_uses_current_state_and_supersedes_old_package(self) -> None:
        cycle_id = self._prepare_cycle()
        with self.app.app_context(), tempfile.TemporaryDirectory() as temp:
            connection = get_db()
            _path, original_payload = create_package_file(
                connection,
                cycle_id=cycle_id,
                manager_pn="111116",
                output_dir=Path(self.app.config["STORAGE_ROOT"]) / "packages",
            )
            original_package_id = original_payload["package"]["package_id"]
            event_count = connection.execute(
                "SELECT COUNT(*) AS count FROM package_events"
            ).fetchone()["count"]
            connection.execute(
                "UPDATE employees SET org_unit = 'abt-aktuell' WHERE pn = '111111'"
            )
            connection.commit()

        page = self.client.get(f"/versand?cycle_id={cycle_id}")
        self.assertIn("START neu", page.get_data(as_text=True))
        response = self.client.post(
            f"/cycles/{cycle_id}/managers/111116/package/recreate"
        )
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.mimetype, "text/html")
        replacement_package_id = response.headers["X-MD-Package-ID"]
        self.assertNotEqual(replacement_package_id, original_package_id)
        self.assertIn(b"abt-aktuell", response.data)
        response.close()

        with self.app.app_context(), tempfile.TemporaryDirectory() as temp:
            connection = get_db()
            self.assertEqual(
                connection.execute(
                    "SELECT COUNT(*) AS count FROM package_events"
                ).fetchone()["count"],
                event_count + 1,
            )
            self.assertEqual(
                connection.execute(
                    "SELECT action FROM audit_log ORDER BY id DESC LIMIT 1"
                ).fetchone()["action"],
                "manager_package_recreated",
            )
            self.assertEqual(
                manager_package_state(
                    connection, cycle_id=cycle_id, manager_pn="111116"
                )["state"],
                "current",
            )
            replacement_candidate = next(
                item for item in dispatch_candidates(connection, cycle_id=cycle_id)
                if item["manager_pn"] == "111116"
            )
            self.assertEqual(replacement_candidate["package_kind"], "start")
            self.assertTrue(replacement_candidate["ready"])
            self.assertIsNone(replacement_candidate["delivery_status"])
            original_payload["package"]["revision"] = 1
            original_payload["package"]["saved_at"] = datetime.now().isoformat()
            old_return = Path(temp) / "alte_arbeitsmappe.html"
            old_return.write_text(render_html(original_payload), encoding="utf-8")
            with self.assertRaisesRegex(ValueError, "neuere START-Datei ersetzt"):
                import_returned_package(
                    connection,
                    path=old_return,
                    original_filename=old_return.name,
                )

    def test_package_prefills_required_scope_and_all_secondary_permissions(self) -> None:
        cycle_id = self._prepare_cycle()
        with self.app.app_context():
            connection = get_db()
            employment = connection.execute(
                """
                SELECT id, last_import_id FROM employment_assignments
                WHERE person_number = '111111' AND assignment_number = '2'
                """
            ).fetchone()
            for fingerprint, permission_for, valid_from, valid_to in (
                ("perm-a", "Vereinsvorstand", "2025-01-01", "2026-12-31"),
                ("perm-b", "Dozententätigkeit", "2025-06-01", "2027-05-31"),
            ):
                connection.execute(
                    """
                    INSERT INTO secondary_activity_permissions (
                        employment_id, valid_from, valid_to, permission,
                        permission_for, source_fingerprint, last_import_id
                    ) VALUES (?, ?, ?, 'Bewilligt', ?, ?, ?)
                    """,
                    (
                        employment["id"], valid_from, valid_to, permission_for,
                        fingerprint, employment["last_import_id"],
                    ),
                )
            connection.commit()
            payload = build_manager_payload(
                connection, cycle_id=cycle_id, manager_pn="111116"
            )
            employee = next(
                item for item in payload["employees"]
                if item["employee"]["pn"] == "111111"
            )
            self.assertEqual(employee["scope"], employee["suggestion"]["scope"])
            self.assertIn("Vereinsvorstand", employee["employee"]["secondary_employment"])
            self.assertIn("Dozententätigkeit", employee["employee"]["secondary_employment"])

    def test_new_sap_import_marks_existing_workbook_for_update(self) -> None:
        cycle_id = self._prepare_cycle()
        with self.app.app_context(), tempfile.TemporaryDirectory() as temp:
            create_package_file(
                get_db(), cycle_id=cycle_id, manager_pn="111116",
                output_dir=Path(temp) / "packages",
            )
        workbook = load_workbook(self.sap_sample)
        worksheet = workbook.active
        headers = {cell.value: cell.column for cell in worksheet[1]}
        for row in range(2, worksheet.max_row + 1):
            if str(worksheet.cell(row, headers["ID_NO_ZERO"]).value) == "111111":
                worksheet.cell(row, headers["OE Bez."], "abt-neu")
        changed = self.sap_sample.with_name("EXPORT_CHANGED.xlsx")
        workbook.save(changed)
        with changed.open("rb") as source:
            response = self.client.post(
                "/imports",
                data={"sap_file": (source, "EXPORT_CHANGED.xlsx")},
                content_type="multipart/form-data",
                follow_redirects=True,
            )
        self.assertEqual(response.status_code, 200)
        with self.app.app_context():
            self.assertEqual(
                manager_package_state(
                    get_db(), cycle_id=cycle_id, manager_pn="111116"
                )["state"],
                "update_required",
            )

    def test_batch_download_contains_required_start_files(self) -> None:
        cycle_id = self._prepare_cycle()
        response = self.client.post(
            f"/cycles/{cycle_id}/packages/batch",
            data={"selection_mode": "all"},
        )
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.mimetype, "application/zip")
        with zipfile.ZipFile(io.BytesIO(response.data)) as archive:
            names = archive.namelist()
        response.close()
        self.assertEqual(len(names), 3)
        self.assertTrue(all(name.startswith("START/") for name in names))

    def test_outlook_draft_is_prepared_and_logged_without_sending(self) -> None:
        cycle_id = self._prepare_cycle()
        with self.app.app_context():
            connection = get_db()
            package_path, _payload = create_package_file(
                connection,
                cycle_id=cycle_id,
                manager_pn="111116",
                output_dir=Path(self.app.config["STORAGE_ROOT"]) / "packages",
            )
            candidates = dispatch_candidates(
                connection, cycle_id=cycle_id, sender_email="md-test@vd.zh.ch"
            )
            candidate = next(
                row for row in candidates if row["manager_pn"] == "111116"
            )
            self.assertTrue(candidate["ready"])
            adapter = FakeOutlookDraftAdapter()
            result = create_outlook_drafts(
                connection,
                cycle_id=cycle_id,
                package_event_ids=[candidate["package_event_id"]],
                sender_email="md-test@vd.zh.ch",
                adapter=adapter,
            )
            self.assertEqual(
                result, {"created": 1, "sent": 0, "failed": 0, "errors": []}
            )
            self.assertEqual(adapter.calls[0]["attachments"], [package_path])
            self.assertEqual(adapter.calls[0]["sender_email"], "md-test@vd.zh.ch")
            delivery = connection.execute(
                "SELECT * FROM mail_deliveries WHERE package_event_id = ?",
                (candidate["package_event_id"],),
            ).fetchone()
            self.assertEqual(delivery["status"], "draft_created")
            self.assertEqual(delivery["encryption_flag_verified"], 1)
            self.assertEqual(delivery["outlook_entry_id"], "draft-1")

        page = self.client.get(f"/versand?cycle_id={cycle_id}")
        body = page.get_data(as_text=True)
        self.assertIn("Entwurf erstellt", body)
        self.assertIn("direkt S/MIME-verschlüsselt versenden", body)
        self.assertIn("E-Mailtext bearbeiten", body)
        self.assertIn("Arbeitsmappen per E-Mail zustellen", body)
        self.assertIn("data-refresh-after-download", body)
        self.assertIn("dispatch.js", body)

    def test_missing_recipient_blocks_outlook_draft(self) -> None:
        cycle_id = self._prepare_cycle()
        with self.app.app_context():
            connection = get_db()
            connection.execute("UPDATE employees SET email = '' WHERE pn = '111116'")
            connection.commit()
            create_package_file(
                connection,
                cycle_id=cycle_id,
                manager_pn="111116",
                output_dir=Path(self.app.config["STORAGE_ROOT"]) / "packages",
            )
            candidate = next(
                row for row in dispatch_candidates(connection, cycle_id=cycle_id)
                if row["manager_pn"] == "111116"
            )
            adapter = FakeOutlookDraftAdapter()
            result = create_outlook_drafts(
                connection,
                cycle_id=cycle_id,
                package_event_ids=[candidate["package_event_id"]],
                adapter=adapter,
            )
            self.assertEqual(result["created"], 0)
            self.assertEqual(result["failed"], 1)
            self.assertEqual(adapter.calls, [])
            self.assertIn("E-Mail-Adresse fehlt", result["errors"][0])

    def test_selected_workbooks_can_be_sent_directly(self) -> None:
        cycle_id = self._prepare_cycle()
        with self.app.app_context():
            connection = get_db()
            create_package_file(
                connection, cycle_id=cycle_id, manager_pn="111116",
                output_dir=Path(self.app.config["STORAGE_ROOT"]) / "packages",
            )
            candidate = next(
                row for row in dispatch_candidates(connection, cycle_id=cycle_id)
                if row["manager_pn"] == "111116"
            )
            adapter = FakeOutlookDraftAdapter()
            result = create_outlook_drafts(
                connection,
                cycle_id=cycle_id,
                package_event_ids=[candidate["package_event_id"]],
                delivery_mode="send",
                subject_template="MD {years} für {manager_name}",
                body_template="Guten Tag {manager_name}\n\n{introduction}",
                adapter=adapter,
            )
            self.assertEqual(result["sent"], 1)
            self.assertEqual(result["created"], 0)
            self.assertEqual(adapter.calls[0]["subject"], "MD 2025/2026 für Rufname6 Nachname6")
            delivery = connection.execute(
                "SELECT status FROM mail_deliveries WHERE package_event_id = ?",
                (candidate["package_event_id"],),
            ).fetchone()
            self.assertEqual(delivery["status"], "sent_confirmed")

    def test_outlook_mailbox_is_resolved_by_display_name(self) -> None:
        class FakeStore:
            DisplayName = "VD-GS HR"

            def GetDefaultFolder(self, folder_kind):
                self.folder_kind = folder_kind
                return "HR-INBOX"

        class FakeRoot:
            Name = "Postfach - VD-GS HR"
            Store = FakeStore()

        class FakeFolders:
            Count = 1

            def Item(self, key):
                if key == 1:
                    return FakeRoot()
                raise RuntimeError("nicht direkt adressierbar")

        class FakeNamespace:
            Folders = FakeFolders()

        self.assertEqual(_resolve_inbox(FakeNamespace(), "VD-GS HR"), "HR-INBOX")

    def test_mail_intake_imports_pdf_and_is_idempotent(self) -> None:
        cycle_id = self._prepare_cycle()
        with self.app.app_context():
            connection = get_db()
            root = Path(self.app.config["STORAGE_ROOT"])
            _package_path, payload = create_package_file(
                connection,
                cycle_id=cycle_id,
                manager_pn="111116",
                output_dir=root / "packages",
            )
            pdf_path = root / "mail-review.pdf"
            write_test_pdf(
                pdf_path,
                "MD-DATENBLOCK ; version=1 ; document=RUECKBLICK ; "
                "case_id=2025-111116-111111 ; "
                f"package_id={payload['package']['package_id']} ; pn=111111 ; ans=2 ; "
                "year=2025 ; scope=review_only ; period_start=2025-01-01 ; "
                "period_end=2025-12-31 ; dialog_date=2026-01-21 ; rating=B ; "
                "agreement=JA ; handwritten_scan_required=NEIN ; no_md_reason=",
            )
            message = InboundMessage(
                entry_id="entry-001",
                internet_message_id="<message-001@example.invalid>",
                sender_email="manager@example.invalid",
                subject="MD Rücklauf",
                received_at=datetime(2026, 1, 22, 10, 0, tzinfo=timezone.utc),
                attachments=(
                    InboundAttachment(1, "Rueckblick_111111.pdf", pdf_path.read_bytes()),
                ),
            )
            moved: list[tuple[str, str]] = []
            first = process_inbound_messages(
                connection,
                messages=[message],
                inbox_dir=root / "mail_inbox",
                pdf_dir=root / "pdf_processed",
                target_folder="12 Mitarbeitenden-Dialog",
                move_message=lambda entry_id, folder: (
                    moved.append((entry_id, folder)) or "entry-001-moved"
                ),
            )
            self.assertEqual(first["ready"], 1)
            self.assertEqual(first["moved"], 1)
            self.assertEqual(moved, [("entry-001", "12 Mitarbeitenden-Dialog")])
            self.assertEqual(first["stored"], 1)
            overview = mail_inbox_overview(connection)
            self.assertEqual(overview["mail_messages"][0]["status"], "moved")
            self.assertEqual(
                overview["mail_messages"][0]["attachments"][0]["status"],
                "imported",
            )
            self.assertEqual(
                connection.execute(
                    "SELECT COUNT(*) AS count FROM official_documents"
                ).fetchone()["count"],
                1,
            )

            repeated = process_inbound_messages(
                connection,
                messages=[message],
                inbox_dir=root / "mail_inbox",
                pdf_dir=root / "pdf_processed",
                target_folder="12 Mitarbeitenden-Dialog",
            )
            self.assertEqual(repeated["duplicates"], 1)
            self.assertEqual(
                connection.execute(
                    "SELECT COUNT(*) AS count FROM inbound_mail_messages"
                ).fetchone()["count"],
                1,
            )

        page = self.client.get(f"/ruecklaeufe?cycle_id={cycle_id}")
        body = page.get_data(as_text=True)
        self.assertIn("MD Rücklauf", body)
        self.assertIn("Verarbeitet und in Outlook verschoben", body)

    def test_feedbacks_are_grouped_per_manager_without_completeness_tracking(self) -> None:
        cycle_id = self._prepare_cycle()
        with self.app.app_context():
            connection = get_db()
            root = Path(self.app.config["STORAGE_ROOT"])
            _package_path, payload = create_package_file(
                connection,
                cycle_id=cycle_id,
                manager_pn="111116",
                output_dir=root / "packages",
            )
            employees = payload["employees"][:2]
            attachments = []
            for index, employee in enumerate(employees, start=1):
                pdf_path = root / f"feedback-{index}.pdf"
                write_test_pdf(
                    pdf_path,
                    "MD-DATENBLOCK ; version=1 ; document=FEEDBACK ; "
                    f"case_id={employee['case_id']} ; "
                    f"package_id={payload['package']['package_id']} ; "
                    f"pn={employee['employee']['pn']} ; manager_pn=111116 ; "
                    "year=2025 ; scope=full",
                )
                attachments.append(
                    InboundAttachment(
                        index,
                        f"{employee['employee']['last_name']}_Feedback_an_VG_111116.pdf",
                        pdf_path.read_bytes(),
                    )
                )
            message = InboundMessage(
                entry_id="entry-feedback",
                internet_message_id="<feedback@example.invalid>",
                sender_email="manager@example.invalid",
                subject="MD Feedbacks gesammelt",
                received_at=datetime(2026, 2, 20, 9, 0, tzinfo=timezone.utc),
                attachments=tuple(attachments),
            )
            result = process_inbound_messages(
                connection,
                messages=[message],
                inbox_dir=root / "mail_inbox",
                pdf_dir=root / "pdf_processed",
                handoff_dir=root / "rpa",
                target_folder="12 Mitarbeitenden-Dialog",
                move_message=lambda entry_id, folder: "entry-feedback-moved",
            )
            self.assertEqual(result["feedback_bundles"], 1)
            self.assertEqual(result["stored"], 2)
            self.assertEqual(result["moved"], 1)
            self.assertEqual(
                connection.execute("SELECT COUNT(*) FROM feedback_submissions").fetchone()[0],
                2,
            )
            bundle = connection.execute("SELECT * FROM feedback_bundles").fetchone()
            self.assertEqual(bundle["manager_pn"], "111116")
            self.assertEqual(bundle["submission_count"], 2)
            self.assertTrue(bundle["filename"].startswith("Feedback_2025_"))
            self.assertEqual(len(PdfReader(bundle["stored_path"]).pages), 3)
            self.assertEqual(
                connection.execute("SELECT COUNT(*) FROM official_documents").fetchone()[0],
                0,
            )

        page = self.client.get(f"/ruecklaeufe?cycle_id={cycle_id}")
        body = page.get_data(as_text=True)
        self.assertIn("Erstellte Sammel-PDFs", body)
        self.assertIn("Fehlende Feedbacks werden nicht überwacht", body)

    def test_mail_without_md_reference_is_ignored_without_storing_attachment(self) -> None:
        self._prepare_cycle()
        with self.app.app_context():
            connection = get_db()
            root = Path(self.app.config["STORAGE_ROOT"])
            message = InboundMessage(
                entry_id="entry-unrelated",
                internet_message_id="<unrelated@example.invalid>",
                sender_email="someone@example.invalid",
                subject="Allgemeine Information",
                received_at=datetime(2026, 1, 22, 12, 0, tzinfo=timezone.utc),
                attachments=(InboundAttachment(1, "Information.txt", b"Kein MD"),),
            )
            result = process_inbound_messages(
                connection,
                messages=[message],
                inbox_dir=root / "mail_inbox",
                pdf_dir=root / "pdf_processed",
                target_folder="12 Mitarbeitenden-Dialog",
            )
            self.assertEqual(result["ignored"], 1)
            self.assertEqual(
                connection.execute("SELECT COUNT(*) FROM inbound_mail_messages").fetchone()[0],
                0,
            )
            self.assertEqual(
                connection.execute("SELECT COUNT(*) FROM inbound_mail_attachments").fetchone()[0],
                0,
            )
            self.assertEqual(
                connection.execute("SELECT COUNT(*) FROM ignored_mail_messages").fetchone()[0],
                1,
            )

    def test_mail_intake_keeps_probation_and_extra_files_for_review(self) -> None:
        cycle_id = self._prepare_cycle()
        with self.app.app_context():
            connection = get_db()
            root = Path(self.app.config["STORAGE_ROOT"])
            _package_path, payload = create_package_file(
                connection,
                cycle_id=cycle_id,
                manager_pn="111116",
                output_dir=root / "packages",
            )
            pdf_path = root / "probation-review.pdf"
            write_test_pdf(
                pdf_path,
                "MD-DATENBLOCK ; version=1 ; document=RUECKBLICK ; "
                "case_id=2025-111116-111111 ; "
                f"package_id={payload['package']['package_id']} ; pn=111111 ; ans=2 ; "
                "year=2025 ; scope=review_only ; period_start=2025-01-01 ; "
                "period_end=2025-12-31 ; dialog_date=2026-01-21 ; rating=B ; "
                "agreement=JA ; handwritten_scan_required=NEIN ; no_md_reason=",
            )
            message = InboundMessage(
                entry_id="entry-probation",
                internet_message_id="<probation@example.invalid>",
                sender_email="manager@example.invalid",
                subject="Rücklauf Probezeit",
                received_at=datetime(2026, 1, 22, 11, 0, tzinfo=timezone.utc),
                attachments=(
                    InboundAttachment(1, "Probezeit_Rueckblick.pdf", pdf_path.read_bytes()),
                    InboundAttachment(2, "Begleitnotiz.txt", b"Bitte beachten"),
                ),
            )
            result = process_inbound_messages(
                connection,
                messages=[message],
                inbox_dir=root / "mail_inbox",
                pdf_dir=root / "pdf_processed",
                target_folder="12 Mitarbeitenden-Dialog",
            )
            self.assertEqual(result["review"], 1)
            stored = connection.execute(
                "SELECT id, status, contains_probation FROM inbound_mail_messages"
            ).fetchone()
            self.assertEqual(stored["status"], "review_required")
            self.assertEqual(stored["contains_probation"], 1)
            cockpit = cockpit_overview(connection, selected_cycle_id=cycle_id)
            self.assertEqual(cockpit["metrics"]["mail_review_required"], 1)
            self.assertTrue(
                any(task["kind"] == "Postfach" for task in cockpit["tasks"])
            )
            attachments = connection.execute(
                "SELECT original_filename, status, stored_path FROM inbound_mail_attachments ORDER BY attachment_index"
            ).fetchall()
            self.assertEqual(
                [(row["original_filename"], row["status"]) for row in attachments],
                [
                    ("Probezeit_Rueckblick.pdf", "imported"),
                    ("Begleitnotiz.txt", "review_required"),
                ],
            )
            self.assertEqual(attachments[1]["stored_path"], "")
            self.assertIn("nicht gespeichert", connection.execute(
                "SELECT error_message FROM inbound_mail_attachments WHERE attachment_index = 2"
            ).fetchone()["error_message"])

            class FakeInboxAdapter:
                def move_message(self, **values):
                    raise AssertionError("Beim Abschluss der Prüfung darf Outlook nicht angesprochen werden.")

            inbox_adapter = FakeInboxAdapter()
            resolved = resolve_inbound_message(
                connection,
                message_id=stored["id"],
                mailbox_name="VD-GS HR",
                target_folder="12 Mitarbeitenden-Dialog",
                adapter=inbox_adapter,
            )
            self.assertEqual(resolved["status"], "review_completed")
            reviewed = connection.execute(
                "SELECT status, review_completed_at FROM inbound_mail_messages WHERE id = ?",
                (stored["id"],),
            ).fetchone()
            self.assertEqual(reviewed["status"], "review_required")
            self.assertTrue(reviewed["review_completed_at"])
            overview = mail_inbox_overview(connection)
            self.assertEqual(overview["mail_messages"][0]["status"], "review_completed")
            cockpit = cockpit_overview(connection, selected_cycle_id=cycle_id)
            self.assertEqual(cockpit["metrics"]["mail_review_required"], 0)

    def test_case_correction_is_reasoned_and_preserves_original_pdf_data(self) -> None:
        cycle_id = self._prepare_cycle()
        with self.app.app_context():
            connection = get_db()
            root = Path(self.app.config["STORAGE_ROOT"])
            _package_path, payload = create_package_file(
                connection, cycle_id=cycle_id, manager_pn="111116",
                output_dir=root / "packages",
            )
            pdf_path = root / "review-original.pdf"
            write_test_pdf(
                pdf_path,
                "MD-DATENBLOCK ; version=1 ; document=RUECKBLICK ; "
                "case_id=2025-111116-111111 ; "
                f"package_id={payload['package']['package_id']} ; pn=111111 ; ans=2 ; "
                "year=2025 ; scope=review_only ; period_start=2025-01-01 ; "
                "period_end=2025-12-31 ; dialog_date=2026-01-21 ; rating=B ; "
                "agreement=JA ; handwritten_scan_required=NEIN ; no_md_reason=",
            )
            imported = import_official_pdf(
                connection,
                path=pdf_path,
                original_filename="Rueckblick_111111.pdf",
                accepted_dir=root / "pdf_processed",
            )
            changed = correct_case_data(
                connection,
                cycle_id=cycle_id,
                case_id="2025-111116-111111",
                values={
                    "official_scope": "review_only",
                    "dialog_date": "2026-01-21",
                    "period_start": "2025-01-01",
                    "period_end": "2025-12-31",
                    "overall_rating_code": "D",
                    "agreement": "NEIN",
                },
                reason="Bewertung gemäss kontrollierter PDF-Prüfung korrigiert.",
                user_id=None,
            )
            self.assertEqual(changed, ["overall_rating_code", "agreement"])
            current = connection.execute(
                "SELECT overall_rating_code, agreement FROM dialog_cases WHERE case_id = ?",
                ("2025-111116-111111",),
            ).fetchone()
            self.assertEqual((current["overall_rating_code"], current["agreement"]), ("D", "NEIN"))
            document = connection.execute(
                "SELECT data_block_json, scan_required FROM official_documents WHERE id = ?",
                (imported["document_id"],),
            ).fetchone()
            self.assertEqual(json.loads(document["data_block_json"])["rating"], "B")
            self.assertEqual(document["scan_required"], 1)
            self.assertEqual(
                connection.execute(
                    "SELECT COUNT(*) AS count FROM case_corrections WHERE case_id = ?",
                    ("2025-111116-111111",),
                ).fetchone()["count"],
                2,
            )
            detail = case_review_detail(
                connection, cycle_id=cycle_id, case_id="2025-111116-111111"
            )
            self.assertFalse(detail["correction_locked"])
            confirm_digital_signature(
                connection,
                document_id=imported["document_id"],
                handoff_dir=root / "dossier_ready",
            )
            with self.assertRaisesRegex(ValueError, "nicht direkt korrigiert"):
                correct_case_data(
                    connection,
                    cycle_id=cycle_id,
                    case_id="2025-111116-111111",
                    values={
                        "official_scope": "review_only",
                        "dialog_date": "2026-01-21",
                        "period_start": "2025-01-01",
                        "period_end": "2025-12-31",
                        "overall_rating_code": "B",
                        "agreement": "JA",
                    },
                    reason="Nachträglicher Änderungsversuch nach Freigabe.",
                    user_id=None,
                )

        page = self.client.get(
            f"/cycles/{cycle_id}/cases/2025-111116-111111"
        )
        body = page.get_data(as_text=True)
        self.assertIn("Einzelfallprüfung", body)
        self.assertIn("Begründete Korrekturen", body)

    def test_corrected_pdf_requires_reason_and_keeps_both_versions(self) -> None:
        cycle_id = self._prepare_cycle()
        with self.app.app_context():
            connection = get_db()
            root = Path(self.app.config["STORAGE_ROOT"])
            _package_path, payload = create_package_file(
                connection, cycle_id=cycle_id, manager_pn="111116",
                output_dir=root / "packages",
            )

            def make_pdf(path: Path, rating: str) -> None:
                write_test_pdf(
                    path,
                    "MD-DATENBLOCK ; version=1 ; document=RUECKBLICK ; "
                    "case_id=2025-111116-111111 ; "
                    f"package_id={payload['package']['package_id']} ; pn=111111 ; ans=2 ; "
                    "year=2025 ; scope=review_only ; period_start=2025-01-01 ; "
                    "period_end=2025-12-31 ; dialog_date=2026-01-21 ; "
                    f"rating={rating} ; agreement=JA ; handwritten_scan_required=NEIN ; "
                    "no_md_reason=",
                )

            first_path = root / "first.pdf"
            make_pdf(first_path, "B")
            first = import_official_pdf(
                connection, path=first_path, original_filename="first.pdf",
                accepted_dir=root / "pdf_processed",
            )
            corrected_path = root / "corrected.pdf"
            make_pdf(corrected_path, "C")
            with self.assertRaisesRegex(ValueError, "Ersetzungsgrund"):
                import_official_pdf(
                    connection, path=corrected_path, original_filename="corrected.pdf",
                    accepted_dir=root / "pdf_processed",
                )
            self.assertTrue(corrected_path.exists())
            second = import_official_pdf(
                connection, path=corrected_path, original_filename="corrected.pdf",
                accepted_dir=root / "pdf_processed",
                replacement_reason="Korrigierte Gesamtbeurteilung nach Rückfrage.",
            )
            versions = connection.execute(
                """
                SELECT id, is_current, replaces_document_id, replacement_reason, replaced_at
                FROM official_documents WHERE case_id = ? ORDER BY id
                """,
                ("2025-111116-111111",),
            ).fetchall()
            self.assertEqual(len(versions), 2)
            self.assertEqual(versions[0]["id"], first["document_id"])
            self.assertEqual(versions[0]["is_current"], 0)
            self.assertTrue(versions[0]["replaced_at"])
            self.assertEqual(versions[1]["id"], second["document_id"])
            self.assertEqual(versions[1]["replaces_document_id"], first["document_id"])
            self.assertIn("Rückfrage", versions[1]["replacement_reason"])

    def test_cycle_creates_dialog_events_and_separate_obligations(self) -> None:
        cycle_id = self._prepare_cycle()
        with self.app.app_context():
            connection = get_db()
            cycle = connection.execute(
                "SELECT review_due_date, outlook_due_date FROM cycles WHERE id = ?",
                (cycle_id,),
            ).fetchone()
            self.assertEqual(cycle["review_due_date"], "2026-01-31")
            self.assertEqual(cycle["outlook_due_date"], "2026-02-28")
            self.assertEqual(
                connection.execute(
                    "SELECT COUNT(*) AS count FROM dialog_events WHERE cycle_id = ?",
                    (cycle_id,),
                ).fetchone()["count"],
                9,
            )
            event = connection.execute(
                """
                SELECT de.id FROM dialog_events de
                WHERE de.legacy_case_id = '2025-111116-111111'
                """
            ).fetchone()
            obligations = connection.execute(
                """
                SELECT document_kind, due_date FROM document_obligations
                WHERE dialog_event_id = ? ORDER BY document_kind
                """,
                (event["id"],),
            ).fetchall()
            self.assertEqual(
                [(row["document_kind"], row["due_date"]) for row in obligations],
                [("outlook", "2026-02-28"), ("review", "2026-01-31")],
            )

    def test_cycle_deadline_extension_updates_open_obligations_and_audit(self) -> None:
        cycle_id = self._prepare_cycle()
        with self.app.app_context():
            connection = get_db()
            result = extend_deadlines(
                connection,
                cycle_id=cycle_id,
                scope="cycle",
                document_kind="review",
                new_due_date="2026-03-31",
                reason="Zusätzliche Bearbeitungszeit gemäss HR-Entscheid",
            )
            self.assertGreater(result["changed"], 0)
            self.assertEqual(
                connection.execute(
                    "SELECT review_due_date FROM cycles WHERE id = ?", (cycle_id,)
                ).fetchone()["review_due_date"],
                "2026-03-31",
            )
            changes = connection.execute(
                "SELECT old_due_date, new_due_date, scope FROM deadline_changes WHERE cycle_id = ?",
                (cycle_id,),
            ).fetchall()
            self.assertEqual(len(changes), result["changed"])
            self.assertTrue(all(row["old_due_date"] == "2026-01-31" for row in changes))
            self.assertTrue(all(row["new_due_date"] == "2026-03-31" for row in changes))
            shortened = extend_deadlines(
                connection,
                cycle_id=cycle_id,
                scope="cycle",
                document_kind="review",
                new_due_date="2026-03-01",
                reason="Begründete Korrektur der zuvor gesetzten Frist",
            )
            self.assertGreater(shortened["changed"], 0)
            with self.assertRaisesRegex(ValueError, "bestehenden Frist"):
                extend_deadlines(
                    connection,
                    cycle_id=cycle_id,
                    scope="cycle",
                    document_kind="review",
                    new_due_date="2026-03-01",
                    reason="Keine tatsächliche Änderung der Frist",
                )

    def test_manager_and_case_deadline_extensions_are_scoped(self) -> None:
        cycle_id = self._prepare_cycle()
        with self.app.app_context():
            connection = get_db()
            manager_result = extend_deadlines(
                connection,
                cycle_id=cycle_id,
                scope="manager",
                manager_pn="111116",
                document_kind="outlook",
                new_due_date="2026-04-15",
                reason="Individuelle Verlängerung der Führungslinie",
            )
            other_due_dates = connection.execute(
                """
                SELECT DISTINCT o.due_date FROM document_obligations o
                JOIN dialog_events de ON de.id = o.dialog_event_id
                JOIN manager_assignments ma ON ma.id = de.manager_assignment_id
                WHERE de.cycle_id = ? AND o.document_kind = 'outlook'
                  AND ma.manager_person_number <> '111116'
                """,
                (cycle_id,),
            ).fetchall()
            self.assertGreater(manager_result["changed"], 0)
            self.assertEqual({row["due_date"] for row in other_due_dates}, {"2026-02-28"})

            case_result = extend_deadlines(
                connection,
                cycle_id=cycle_id,
                scope="case",
                case_id="2025-111116-111111",
                document_kind="review",
                new_due_date="2026-05-01",
                reason="Einzelfallverlängerung nach dokumentierter Rückfrage",
            )
            self.assertEqual(case_result["changed"], 1)

    def test_reminder_draft_is_encrypted_logged_and_not_duplicated(self) -> None:
        cycle_id = self._prepare_cycle()
        with self.app.app_context():
            connection = get_db()
            candidates = reminder_candidates(
                connection, cycle_id=cycle_id, today=date(2026, 3, 1)
            )
            candidate = next(row for row in candidates if row["recipient_email"])
            adapter = FakeOutlookDraftAdapter()
            result = create_reminder_drafts(
                connection,
                cycle_id=cycle_id,
                manager_pns=[candidate["manager_pn"]],
                adapter=adapter,
                today=date(2026, 3, 1),
                created_at=datetime(2026, 3, 1, 9, 0, tzinfo=timezone.utc),
            )
            self.assertEqual(result["created"], 1)
            self.assertEqual(adapter.calls[0]["attachments"], [])
            reminder = connection.execute(
                "SELECT * FROM reminders WHERE cycle_id = ? AND manager_pn = ?",
                (cycle_id, candidate["manager_pn"]),
            ).fetchone()
            self.assertEqual(reminder["status"], "draft_created")
            self.assertEqual(reminder["encryption_flag_verified"], 1)
            snapshots = connection.execute(
                "SELECT COUNT(*) AS count FROM reminder_obligations WHERE reminder_id = ?",
                (reminder["id"],),
            ).fetchone()["count"]
            self.assertEqual(snapshots, candidate["overdue_count"])
            with self.assertRaisesRegex(ValueError, "offener Entwurf"):
                create_reminder_drafts(
                    connection,
                    cycle_id=cycle_id,
                    manager_pns=[candidate["manager_pn"]],
                    adapter=adapter,
                    today=date(2026, 3, 1),
                )
            confirmed = confirm_reminder_sent(connection, reminder_id=reminder["id"])
            self.assertEqual(confirmed["status"], "sent_confirmed")

    def test_dialog_page_offers_deadlines_and_reminder_preview(self) -> None:
        cycle_id = self._prepare_cycle()
        response = self.client.get(f"/dialoge?cycle_id={cycle_id}")
        self.assertEqual(response.status_code, 200)
        page = response.get_data(as_text=True)
        self.assertIn("Fristen des regulären Durchlaufs", page)
        self.assertIn("Rückblick:</strong> 2026-01-31", page)
        self.assertIn("Ausblick:</strong> 2026-02-28", page)
        self.assertIn("maximal 50 pro Seite", page)
        self.assertIn("Erinnerungsentwürfe", page)
        self.assertIn("S/MIME-markierte Outlook-Entwürfe", page)

    def test_analytics_filters_aggregates_and_exports_without_small_groups(self) -> None:
        cycle_id = self._prepare_cycle()
        with self.app.app_context():
            connection = get_db()
            cases = connection.execute(
                """
                SELECT dc.case_id, e.org_unit, dc.employee_pn
                FROM dialog_cases dc
                JOIN dialog_events de ON de.legacy_case_id = dc.case_id
                JOIN employees e ON e.pn = dc.employee_pn
                WHERE dc.cycle_id = ? AND de.required_scope = 'full'
                ORDER BY dc.id LIMIT 5
                """,
                (cycle_id,),
            ).fetchall()
            self.assertEqual(len(cases), 5)
            ratings = ["A", "B", "B", "C", "D"]
            for index, case in enumerate(cases):
                payload = {
                    "scope": "full",
                    "review": {
                        "dialog_date": f"2026-01-{10 + index:02d}",
                        "overall_rating": ratings[index],
                        "competencies": [{"competency": "Kooperationsfähigkeit"}],
                    },
                    "previous_development_goals": [],
                    "outlook": {
                        "dialog_date": f"2026-02-{10 + index:02d}",
                        "development_goals": [
                            {"competency": "Entwicklungsfähigkeit"}
                        ],
                    },
                }
                connection.execute(
                    """
                    UPDATE dialog_cases SET data_json = ?, official_scope = 'full',
                        dialog_date = ?, overall_rating_code = ?, status = 'vollstaendig'
                    WHERE case_id = ?
                    """,
                    (
                        json.dumps(payload, ensure_ascii=False),
                        payload["review"]["dialog_date"], ratings[index], case["case_id"],
                    ),
                )
            connection.commit()

            data = analytics_data(
                connection, cycle_id=cycle_id, minimum_group_size=5
            )
            self.assertEqual(data["rating_group_size"], 5)
            self.assertFalse(data["ratings_suppressed"])
            self.assertEqual(
                {row["rating"]: row["count"] for row in data["ratings"]}["B"], 2
            )
            self.assertEqual(len(data["timings"]), 10)
            competency_counts = {
                (row["competency"], row["usage"]): row["count"]
                for row in data["competencies"]
            }
            self.assertEqual(
                competency_counts[("Kooperationsfähigkeit", "Im Rückblick thematisiert")],
                5,
            )
            self.assertEqual(
                competency_counts[("Entwicklungsfähigkeit", "Als Entwicklungsziel vereinbart")],
                5,
            )
            small = analytics_data(
                connection,
                cycle_id=cycle_id,
                org_unit=cases[0]["org_unit"],
                minimum_group_size=5,
            )
            self.assertTrue(small["ratings_suppressed"])

        response = self.client.post(
            "/auswertungen/export",
            data={"report": "ratings", "cycle_id": str(cycle_id)},
        )
        self.assertEqual(response.status_code, 200)
        self.assertTrue(response.data.startswith(b"\xef\xbb\xbf"))
        self.assertIn(b"Gesamtbeurteilung;Anzahl;Anteil Prozent", response.data)
        page = self.client.get(f"/auswertungen?cycle_id={cycle_id}").get_data(as_text=True)
        self.assertIn("Gesprächszeitpunkte", page)
        self.assertIn("Kooperationsfähigkeit", page)
        self.assertIn("<progress", page)
        self.assertIn("table_controls.js", page)

    def test_workbook_pdf_data_block_contains_structured_competencies(self) -> None:
        template = (PROJECT_ROOT / "prototype" / "html_dialog" / "template.html").read_text(
            encoding="utf-8"
        )
        self.assertIn("review_competencies=", template)
        self.assertIn("development_competencies=", template)

    def test_manual_dialog_event_can_be_added(self) -> None:
        cycle_id = self._prepare_cycle()
        response = self.client.post(
            "/dialog-events",
            data={
                "employee_pn": "111111",
                "assignment_number": "2",
                "manager_pn": "111118",
                "event_type": "probation_review",
                "required_scope": "review_only",
                "review_year": "2025",
                "period_start": "2025-01-01",
                "period_end": "2025-03-31",
                "dialog_date": "2025-03-20",
                "review_due_date": "2025-04-15",
                "outlook_due_date": "",
                "cycle_id": str(cycle_id),
                "reason": "Probezeitfall gemäss HR-Abklärung",
            },
            follow_redirects=True,
        )
        self.assertEqual(response.status_code, 200)
        self.assertIn("Rückblick Probezeit", response.get_data(as_text=True))
        with self.app.app_context():
            connection = get_db()
            event = connection.execute(
                """
                SELECT * FROM dialog_events
                WHERE source = 'manual' AND event_type = 'probation_review'
                """
            ).fetchone()
            self.assertIsNotNone(event)
            self.assertEqual(event["dialog_date"], "2025-03-20")
            obligation = connection.execute(
                "SELECT * FROM document_obligations WHERE dialog_event_id = ?",
                (event["id"],),
            ).fetchone()
            self.assertEqual(obligation["document_kind"], "review")
            self.assertEqual(obligation["due_date"], "2025-04-15")
            payload = build_manager_payload(
                connection, cycle_id=cycle_id, manager_pn="111118"
            )
            manual_case = next(
                item for item in payload["employees"]
                if item["case_id"] == event["event_id"]
            )
            self.assertEqual(manual_case["dialog_type"], "probation")
            self.assertEqual(manual_case["suggestion"]["scope"], "review_only")
            set_leading_sap_event(connection, event["id"])

    def test_sap_import_cycle_and_manager_counts(self) -> None:
        cycle_id = self._prepare_cycle()
        with self.app.app_context():
            connection = get_db()
            cycle, managers = cycle_overview(connection, cycle_id)
            self.assertEqual(cycle["review_year"], 2025)
            counts = {row["manager_pn"]: row["case_count"] for row in managers}
            self.assertEqual(counts["111116"], 3)
            self.assertEqual(sum(counts.values()), 9)
            assignment = connection.execute(
                "SELECT employment_assignment FROM employees WHERE pn = '111111'"
            ).fetchone()["employment_assignment"]
            self.assertEqual(assignment, "2")
        response = self.client.get(f"/cycles/{cycle_id}/managers/111116")
        self.assertEqual(response.status_code, 200)
        self.assertIn("Rufname1 Nachname1", response.get_data(as_text=True))

    def test_review_period_is_clipped_to_entry_and_exit(self) -> None:
        cycle_id = self._prepare_cycle()
        with self.app.app_context():
            connection = get_db()
            exit_payload = build_manager_payload(
                connection, cycle_id=cycle_id, manager_pn="111116"
            )
            exiting = next(
                item for item in exit_payload["employees"] if item["employee"]["pn"] == "111112"
            )
            self.assertEqual(exiting["period_start"], "2025-01-01")
            self.assertEqual(exiting["period_end"], "2025-10-31")
            self.assertEqual(exiting["suggestion"]["scope"], "review_only")
            self.assertIn(
                "Auf Wunsch des Mitarbeitenden kann auf den Rückblick verzichtet werden.",
                exiting["suggestion"]["reason"],
            )

            entry_payload = build_manager_payload(
                connection, cycle_id=cycle_id, manager_pn="111118"
            )
            entrant = next(
                item for item in entry_payload["employees"] if item["employee"]["pn"] == "111114"
            )
            self.assertEqual(entrant["period_start"], "2025-09-01")
            self.assertEqual(entrant["period_end"], "2025-12-31")

    def test_sap_export_uses_only_required_columns_and_clipped_period(self) -> None:
        cycle_id = self._prepare_cycle()
        with self.app.app_context(), tempfile.TemporaryDirectory() as temp:
            connection = get_db()
            payload = build_manager_payload(
                connection, cycle_id=cycle_id, manager_pn="111116"
            )
            employee = next(
                item for item in payload["employees"] if item["employee"]["pn"] == "111112"
            )
            employee["scope"] = "review_only"
            employee["scope_reason"] = "Austritt"
            employee["dialog_date"] = "2026-01-21"
            employee["review"]["overall_rating"] = "D – genügend"
            employee["review"]["agreement"] = "Nein"
            connection.execute(
                """
                UPDATE dialog_cases
                SET status = 'vollstaendig', data_json = ?
                WHERE cycle_id = ? AND employee_pn = '111112' AND manager_pn = '111116'
                """,
                (json.dumps(employee, ensure_ascii=False), cycle_id),
            )
            event = connection.execute(
                "SELECT id FROM dialog_events WHERE legacy_case_id = '2025-111116-111112'"
            ).fetchone()
            set_leading_sap_event(connection, event["id"])
            connection.commit()
            self.assertEqual(
                cockpit_overview(connection, selected_cycle_id=cycle_id)["metrics"][
                    "sap_exports_pending"
                ],
                1,
            )
            path = create_sap_upload_file(
                connection,
                cycle_id=cycle_id,
                template_path=PROJECT_ROOT / "templates" / "massenupload.xlsx",
                output_dir=Path(temp),
                created_at=datetime(2026, 2, 2, 10, 30, tzinfo=timezone.utc),
            )

            worksheet = load_workbook(path).active
            self.assertEqual(worksheet.max_row, 2)
            self.assertEqual(worksheet["A2"].value, 111112)
            self.assertEqual(worksheet["B2"].value, 1)
            self.assertEqual(worksheet["C2"].value.date().isoformat(), "2025-01-01")
            self.assertEqual(worksheet["D2"].value.date().isoformat(), "2025-10-31")
            self.assertEqual(worksheet["E2"].value, 1)
            self.assertEqual(worksheet["F2"].value.date().isoformat(), "2026-01-21")
            self.assertEqual(worksheet["G2"].value.date().isoformat(), "2025-01-01")
            self.assertEqual(worksheet["H2"].value.date().isoformat(), "2025-10-31")
            self.assertEqual(worksheet["I2"].value, "D")
            self.assertTrue(all(worksheet.cell(2, column).value is None for column in range(10, 15)))

            batch = connection.execute(
                "SELECT * FROM sap_export_batches WHERE cycle_id = ?", (cycle_id,)
            ).fetchone()
            self.assertIsNotNone(batch)
            self.assertEqual(batch["row_count"], 1)
            self.assertEqual(batch["filename"], path.name)
            exported_row = connection.execute(
                "SELECT * FROM sap_export_rows WHERE batch_id = ?", (batch["id"],)
            ).fetchone()
            self.assertEqual(exported_row["employee_pn"], "111112")
            self.assertEqual(exported_row["overall_rating"], "D")
            self.assertTrue(exported_row["source_version"])
            self.assertEqual(
                cockpit_overview(connection, selected_cycle_id=cycle_id)["metrics"][
                    "sap_exports_pending"
                ],
                0,
            )

            with self.assertRaisesRegex(ValueError, "keine neuen"):
                create_sap_upload_file(
                    connection,
                    cycle_id=cycle_id,
                    template_path=PROJECT_ROOT / "templates" / "massenupload.xlsx",
                    output_dir=Path(temp),
                    created_at=datetime(2026, 2, 2, 10, 31, tzinfo=timezone.utc),
                )

            response = self.client.post(
                f"/sap-export-batches/{batch['id']}/download"
            )
            self.assertEqual(response.status_code, 200)
            self.assertIn(
                path.name,
                response.headers["Content-Disposition"],
            )
            response.close()
            page = self.client.get("/sap-export").get_data(as_text=True)
            self.assertIn("Schutz vor Doppelverarbeitung", page)
            self.assertIn(path.name, page)

    def test_pdf_only_return_is_read_and_staged_for_dossier_handoff(self) -> None:
        cycle_id = self._prepare_cycle()
        with self.app.app_context(), tempfile.TemporaryDirectory() as temp:
            connection = get_db()
            root = Path(temp)
            _package_path, payload = create_package_file(
                connection,
                cycle_id=cycle_id,
                manager_pn="111116",
                output_dir=root / "packages",
            )
            pdf_path = root / "review.pdf"
            write_test_pdf(
                pdf_path,
                "MD-DATENBLOCK ; version=1 ; document=RUECKBLICK ; "
                "case_id=2025-111116-111111 ; "
                f"package_id={payload['package']['package_id']} ; pn=111111 ; ans=2 ; "
                "year=2025 ; scope=review_only ; period_start=2025-01-01 ; "
                "period_end=2025-12-31 ; dialog_date=2026-01-21 ; rating=B ; "
                "agreement=JA ; handwritten_scan_required=NEIN ; no_md_reason=",
            )
            imported = import_official_pdf(
                connection,
                path=pdf_path,
                original_filename="Rueckblick_2025_Nachname1_Rufname1_111111.pdf",
                accepted_dir=root / "processed",
                handoff_dir=root / "dossier_ready",
            )
            self.assertFalse(pdf_path.exists())
            case = connection.execute(
                "SELECT * FROM dialog_cases WHERE case_id = '2025-111116-111111'"
            ).fetchone()
            self.assertEqual(case["official_scope"], "review_only")
            self.assertEqual(case["overall_rating_code"], "B")
            self.assertEqual(case["status"], "vollstaendig")
            staged = connection.execute(
                "SELECT handoff_filename FROM official_documents WHERE id = ?",
                (imported["document_id"],),
            ).fetchone()
            self.assertEqual(
                staged["handoff_filename"],
                "Rueckblick_2025_Nachname1_Rufname1_111111.pdf",
            )
            self.assertTrue((root / "dossier_ready" / staged["handoff_filename"]).exists())
            status = connection.execute(
                "SELECT status FROM dialog_cases WHERE case_id = '2025-111116-111111'"
            ).fetchone()["status"]
            self.assertEqual(status, "vollstaendig")
            obligation = connection.execute(
                "SELECT status FROM document_obligations WHERE fulfilled_document_id = ?",
                (imported["document_id"],),
            ).fetchone()
            self.assertEqual(obligation["status"], "complete")

    def test_required_scan_replaces_digital_pdf_for_dossier_handoff(self) -> None:
        cycle_id = self._prepare_cycle()
        with self.app.app_context(), tempfile.TemporaryDirectory() as temp:
            connection = get_db()
            root = Path(temp)
            _package_path, payload = create_package_file(
                connection,
                cycle_id=cycle_id,
                manager_pn="111116",
                output_dir=root / "packages",
            )
            pdf_path = root / "review-d.pdf"
            write_test_pdf(
                pdf_path,
                "MD-DATENBLOCK ; version=1 ; document=RUECKBLICK ; "
                "case_id=2025-111116-111111 ; "
                f"package_id={payload['package']['package_id']} ; pn=111111 ; ans=2 ; "
                "year=2025 ; scope=review_only ; period_start=2025-01-01 ; "
                "period_end=2025-12-31 ; dialog_date=2026-01-21 ; rating=D ; "
                "agreement=NEIN ; handwritten_scan_required=JA ; no_md_reason=",
            )
            imported = import_official_pdf(
                connection,
                path=pdf_path,
                original_filename="Rueckblick_2025_Nachname1_Rufname1_111111.pdf",
                accepted_dir=root / "processed",
                handoff_dir=root / "dossier_ready",
            )
            self.assertTrue(imported["scan_required"])
            self.assertEqual(list((root / "dossier_ready").glob("*.pdf")), [])
            obligation = connection.execute(
                "SELECT status FROM document_obligations WHERE fulfilled_document_id = ?",
                (imported["document_id"],),
            ).fetchone()
            self.assertEqual(obligation["status"], "scan_pending")

            scan_path = root / "scan.pdf"
            write_test_pdf(scan_path, "Handschriftlich unterzeichnet")
            scan = import_handwritten_scan(
                connection,
                case_id="2025-111116-111111",
                document_kind="review",
                path=scan_path,
                original_filename="scan.pdf",
                handoff_dir=root / "dossier_ready",
            )
            self.assertEqual(
                scan["handoff_filename"],
                "Rueckblick_2025_Nachname1_Rufname1_HANDSCAN_111111.pdf",
            )
            handoff_files = [path.name for path in (root / "dossier_ready").glob("*.pdf")]
            self.assertEqual(handoff_files, [scan["handoff_filename"]])
            digital = connection.execute(
                "SELECT * FROM official_documents WHERE id = ?", (imported["document_id"],)
            ).fetchone()
            self.assertTrue(Path(digital["stored_path"]).exists())
            self.assertNotEqual(Path(digital["stored_path"]).parent, root / "dossier_ready")
            obligation = connection.execute(
                "SELECT status, fulfilled_document_id FROM document_obligations WHERE id = ?",
                (digital["document_obligation_id"],),
            ).fetchone()
            self.assertEqual(obligation["status"], "complete")
            self.assertEqual(obligation["fulfilled_document_id"], scan["document_id"])

    def test_no_md_pdf_closes_case_without_dossier_handoff(self) -> None:
        cycle_id = self._prepare_cycle()
        with self.app.app_context(), tempfile.TemporaryDirectory() as temp:
            connection = get_db()
            root = Path(temp)
            _package_path, payload = create_package_file(
                connection,
                cycle_id=cycle_id,
                manager_pn="111116",
                output_dir=root / "packages",
            )
            pdf_path = root / "no-md.pdf"
            write_test_pdf(
                pdf_path,
                "MD-DATENBLOCK ; version=1 ; document=KEIN_MD ; "
                "case_id=2025-111116-111112 ; "
                f"package_id={payload['package']['package_id']} ; pn=111112 ; ans=1 ; "
                "year=2025 ; scope=none ; period_start=2025-01-01 ; "
                "period_end=2025-10-31 ; dialog_date= ; rating= ; agreement= ; "
                "handwritten_scan_required=NEIN ; no_md_reason=Austritt",
            )
            imported = import_official_pdf(
                connection,
                path=pdf_path,
                original_filename="Kein_MD_2025_Nachname2_Rufname2_111112.pdf",
                accepted_dir=root / "processed",
            )
            self.assertEqual(imported["document_kind"], "no_md")
            case = connection.execute(
                "SELECT status, official_scope FROM dialog_cases WHERE case_id = '2025-111116-111112'"
            ).fetchone()
            self.assertEqual(case["status"], "kein_md")
            self.assertEqual(case["official_scope"], "none")
            self.assertFalse((root / "dossier_ready").exists())

    def test_package_download_and_saved_return_roundtrip(self) -> None:
        cycle_id = self._prepare_cycle()
        response = self.client.post(f"/cycles/{cycle_id}/managers/111116/package")
        self.assertEqual(response.status_code, 200)
        self.assertIn(
            "MD_Dialog_2025_2026_Nachname6_Rufname6_111116_START.html",
            response.headers["Content-Disposition"],
        )

        with tempfile.TemporaryDirectory() as temp:
            start_path = Path(temp) / "start.html"
            start_path.write_bytes(response.data)
            payload = extract_payload(start_path)
            start_html = start_path.read_text(encoding="utf-8")
        response.close()
        preparation_template = payload["configuration"]["preparation_template"]
        self.assertIn("Gesprächsvorbereitung als PDF", preparation_template)
        self.assertIn("Feedback als PDF", preparation_template)
        self.assertIn("Teilkompetenzen und Formulierungshilfen", preparation_template)
        self.assertIn('data-action="print-preparation">Gesprächsvorbereitung als PDF', preparation_template)
        self.assertIn("Ungespeicherte Angaben", preparation_template)
        self.assertIn("Schliesse diese bisherige Datei", preparation_template)
        self.assertIn("const markDirty", preparation_template)
        self.assertNotIn("HR fordert fehlende Feedbacks nicht nach.", preparation_template)
        self.assertNotIn("gesammelt und S/MIME-verschlüsselt", preparation_template)
        self.assertIn("Die Führungskräfte sammeln", self.client.get(
            f"/ruecklaeufe?cycle_id={cycle_id}"
        ).get_data(as_text=True))
        self.assertIn("2 · Vorbereitung MA", start_html)
        self.assertIn("selectedStepByCase[employee.case_id] || 'basics'", start_html)
        self.assertIn("selectedStep === 'preparation'", start_html)
        self.assertIn('<div class="suggestion warning"><strong>Feedback-Rücklauf:', start_html)
        self.assertNotIn("HR fordert fehlende Feedbacks nicht nach.</div>", start_html)
        self.assertIn("if (hasReview)", preparation_template)
        self.assertIn("if (hasOutlook)", preparation_template)
        payload["package"]["revision"] = 1
        payload["package"]["saved_at"] = "2025-09-17T12:30:00+00:00"
        returned_html = render_html(payload).encode("utf-8")
        response = self.client.post(
            "/returns",
            data={
                "return_file": (
                    io.BytesIO(returned_html),
                    "MD_Dialog_2025_2026_Nachname6_Rufname6_111116_BEARBEITET_v01_20250917_1230.html",
                )
            },
            content_type="multipart/form-data",
            follow_redirects=True,
        )
        self.assertEqual(response.status_code, 200)
        self.assertIn("Rücklauf von Rufname6 Nachname6", response.get_data(as_text=True))

        duplicate_response = self.client.post(
            "/returns",
            data={
                "return_file": (
                    io.BytesIO(returned_html),
                    "MD_Dialog_2025_2026_Nachname6_Rufname6_111116_BEARBEITET_v01_20250917_1230.html",
                )
            },
            content_type="multipart/form-data",
            follow_redirects=True,
        )
        self.assertIn("wurde bereits importiert oder ist veraltet", duplicate_response.get_data(as_text=True))

        with self.app.app_context():
            connection = get_db()
            events = connection.execute(
                "SELECT direction FROM package_events ORDER BY id"
            ).fetchall()
            self.assertEqual([row["direction"] for row in events], ["versand", "ruecklauf"])
            _cycle, managers = cycle_overview(connection, cycle_id)
            manager = next(row for row in managers if row["manager_pn"] == "111116")
            self.assertEqual(manager["case_count"], 3)

    def test_start_file_cannot_be_imported_as_return(self) -> None:
        cycle_id = self._prepare_cycle()
        package = self.client.post(f"/cycles/{cycle_id}/managers/111116/package")
        package_data = package.data
        package.close()
        response = self.client.post(
            "/returns",
            data={"return_file": (io.BytesIO(package_data), "START.html")},
            content_type="multipart/form-data",
            follow_redirects=True,
        )
        self.assertIn("Die Datei ist noch eine Startdatei", response.get_data(as_text=True))


if __name__ == "__main__":
    unittest.main()
